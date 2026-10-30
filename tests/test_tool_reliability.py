import asyncio
import json
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from pathlib import Path
from threading import Barrier
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from sqlalchemy import create_engine, inspect, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker

from app.agents.harness import MindBridgeAgentHarness
from app.core.config import Settings
from app.core.database import Base
from app.core.tool_job_migration import upgrade_tool_jobs
from app.models.entities import (
    AlertRecord, CaseNote, ChatMessage, ChatSession, DeadLetterRecord, ExcelRecord,
    PsychologicalReport, RiskCase, ToolAuditRecord, ToolJob, UserAccount,
)
from app.schemas.dtos import ChatRequest
from app.services.chat import ChatService
from app.services.mcp_client import McpToolError, MindBridgeMcpToolClient
from app.services.tool_execution import ToolExecutionError, ToolExecutionService
from app.services.tool_governance import ToolAuthorizationError
from app.services.tool_queue import ToolQueueService, ToolQueueWorker
from app.services.tools import ToolOrchestrationService


class ToolRuntimeFixture(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.engine = create_engine(
            f"sqlite:///{self.root / 'test.sqlite3'}", connect_args={"check_same_thread": False, "timeout": 30},
        )
        self.addCleanup(self.engine.dispose)
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine, autoflush=False)
        factory = patch("app.services.tool_queue.SessionLocal", self.sessions)
        factory.start()
        self.addCleanup(factory.stop)
        memory = patch("app.services.memory.RedisShortTermMemoryStore._connect", return_value=None)
        memory.start()
        self.addCleanup(memory.stop)
        self.settings = Settings(
            _env_file=None, database_url=str(self.engine.url), ai_provider="mock",
            knowledge_vector_enabled=False, tool_queue_enabled=True,
            excel_path=str(self.root / 'ledger.xlsx'), alert_email_delivery_mode="log",
            tool_queue_retry_delay_seconds=0.01,
            privacy_key_file=str(self.root/'fixture.key'),
        )
        from app.core.privacy_crypto import cipher
        crypto=patch('app.core.privacy_crypto.get_settings',return_value=self.settings)
        crypto.start();self.addCleanup(crypto.stop);cipher(self.settings)
        self.db = self.sessions()
        self.addCleanup(self.db.close)
        self.user = UserAccount(username="student", display_name="Test Student", password_hash="unused")
        admin = UserAccount(username="admin", display_name="Test Admin", password_hash="unused")
        admin.roles = {"ROLE_ADMIN", "ROLE_USER"}
        self.db.add_all([self.user, admin])
        self.db.flush()
        self.session = ChatSession(user_id=self.user.id, public_id="test-session", title="Synthetic test")
        self.db.add(self.session)
        self.db.commit()
        self.worker = ToolQueueWorker(self.settings)
        self.addCleanup(self.worker.stop)

    def report(self, risk="HIGH"):
        row = PsychologicalReport(
            user_id=self.user.id, session_id=self.session.id, content="synthetic report",
            intent="RISK" if risk == "HIGH" else "CONSULT", emotion="HIGH_RISK" if risk == "HIGH" else "ANXIETY",
            emotion_score=4.0 if risk == "HIGH" else 2.0, risk_level=risk,
            confidence=0.9, summary="synthetic test",
        )
        self.db.add(row)
        self.db.commit()
        return row

    def rows(self, model):
        with self.sessions() as db:
            return db.query(model).order_by(model.id).all()

    def drain(self, jobs):
        for job in jobs:
            self.worker._run_job(job.id)

    def retry_now(self, job_id):
        with self.sessions() as db:
            db.get(ToolJob, job_id).run_after = datetime.utcnow() - timedelta(seconds=1)
            db.commit()

    def events(self, chunks):
        return [json.loads(line[6:]) for chunk in chunks for line in chunk.splitlines() if line.startswith("data: ")]


class OutboxReliabilityTests(ToolRuntimeFixture):
    def test_report_and_all_tasks_rollback_when_enqueue_fails_halfway(self):
        original = ToolQueueService._find_or_create

        def fail_case(service, kind, *args):
            if kind == "CASE_CREATE":
                raise RuntimeError("injected enqueue failure")
            return original(service, kind, *args)

        with patch.object(ToolQueueService, "_find_or_create", fail_case):
            with self.assertRaisesRegex(RuntimeError, "injected"):
                MindBridgeAgentHarness(self.db, self.settings).run(
                    self.user, ChatRequest(message="我不想活了。", sessionId=self.session.public_id),
                )
        self.assertEqual(self.rows(PsychologicalReport), [])
        self.assertEqual(self.rows(ToolJob), [])

    def test_client_disconnect_after_meta_does_not_cancel_delivery(self):
        async def disconnect():
            stream = ChatService(self.db, self.settings).stream_chat(
                self.user, ChatRequest(message="我不想活了。", sessionId=self.session.public_id),
            )
            self.assertIn("event: meta", await anext(stream))
            self.assertEqual(len(self.rows(ToolJob)), 3)
            await stream.aclose()

        asyncio.run(disconnect())
        self.drain(self.rows(ToolJob))
        self.assertEqual({job.status for job in self.rows(ToolJob)}, {"SUCCESS"})
        self.assertEqual(len(self.rows(AlertRecord)), 1)
        self.assertEqual(len(self.rows(ChatMessage)), 1)

    def test_model_failure_and_empty_reply_emit_error_but_keep_tasks(self):
        for empty in [False, True]:
            with self.subTest(empty=empty):
                service = ChatService(self.db, self.settings)

                async def broken(messages):
                    self.assertEqual(len(self.rows(ToolJob)), 3 if not empty else 6)
                    if not empty:
                        raise ConnectionError("injected model outage")
                    if False:
                        yield ""

                async def collect():
                    return [chunk async for chunk in service.stream_chat(
                        self.user, ChatRequest(message="我不想活了。", sessionId=self.session.public_id),
                    )]

                with patch.object(service.ai, "stream", broken):
                    events = self.events(asyncio.run(collect()))
                self.assertIn("error", [event["type"] for event in events])
                self.assertNotIn("done", [event["type"] for event in events])
                self.assertTrue(any(event.get("content") for event in events if event["type"] == "token"))
                self.assertNotIn("injected", json.dumps(events))
        self.drain(self.rows(ToolJob))
        self.assertEqual(len(self.rows(AlertRecord)), 2)

    def test_repeated_enqueue_uses_persisted_risk_and_keeps_dead_task(self):
        report = self.report("LOW")
        jobs = ToolQueueService(self.db, self.settings).enqueue_report(report.id, "HIGH")
        self.assertEqual([job.kind for job in jobs], ["EXCEL_REPORT"])
        jobs[0].status = "DEAD"
        self.db.commit()
        again = ToolQueueService(self.db, self.settings).enqueue_report(report.id, "HIGH")
        self.assertEqual([job.id for job in again], [jobs[0].id])
        self.assertEqual(len(self.rows(ToolJob)), 1)

    def test_two_workers_cannot_claim_the_same_task(self):
        job = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)[0]
        # Pass only the scalar ID across threads; an expired ORM instance would
        # otherwise make both threads use the fixture's non-thread-safe Session.
        job_id = job.id
        barrier = Barrier(2)

        def claim():
            barrier.wait()
            return self.worker._claim_job(job_id)

        with ThreadPoolExecutor(max_workers=2) as pool:
            tokens = list(pool.map(lambda _: claim(), range(2)))
        self.assertEqual(sum(token is not None for token in tokens), 1)

    def test_expired_lease_recovers_and_stale_worker_cannot_execute(self):
        job = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)[0]
        old_token = self.worker._claim_job(job.id)
        self.worker._recover_running_jobs()
        self.assertEqual(self.rows(ToolJob)[0].claim_token, old_token)
        with self.sessions() as db:
            db.get(ToolJob, job.id).lease_expires_at = datetime.utcnow() - timedelta(seconds=1)
            db.commit()
        restarted = ToolQueueWorker(self.settings)
        self.addCleanup(restarted.stop)
        restarted._recover_running_jobs()
        self.worker._run_job(job.id, old_token)
        self.assertEqual(self.rows(ExcelRecord), [])
        restarted._run_job(job.id)
        self.assertEqual(len(self.rows(ExcelRecord)), 1)
        self.assertEqual(self.rows(ToolJob)[0].status, "SUCCESS")

    def test_lost_success_reply_is_recovered_from_committed_audit(self):
        job = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)[0]
        token = self.worker._claim_job(job.id)
        with self.sessions() as db:
            ToolExecutionService(db, self.settings).execute(
                job.kind, report_id=job.report_id, job_id=job.id, claim_token=token,
            )
            db.get(ToolJob, job.id).lease_expires_at = datetime.utcnow() - timedelta(seconds=1)
            db.commit()
        self.worker._recover_running_jobs()
        with patch.object(ToolExecutionService, "execute", side_effect=AssertionError("must not replay")):
            self.worker._run_job(job.id)
        self.assertEqual(self.rows(ToolJob)[0].status, "SUCCESS")
        self.assertEqual(len(self.rows(ExcelRecord)), 1)

    def test_lease_expiring_before_execution_requeues_without_consuming_attempt(self):
        job = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)[0]
        token = self.worker._claim_job(job.id)
        with self.sessions() as db:
            db.get(ToolJob, job.id).lease_expires_at = datetime.utcnow() - timedelta(seconds=1)
            db.commit()
        self.worker._run_job(job.id, token)
        stored = self.rows(ToolJob)[0]
        self.assertEqual(stored.status, "PENDING")
        self.assertEqual(stored.attempts, 0)
        self.assertEqual(self.rows(DeadLetterRecord), [])

    def test_ledger_retries_after_db_rollback_do_not_duplicate_file_row(self):
        report = self.report()
        tools = ToolOrchestrationService(self.db, self.settings, autocommit=False)
        tools.write_excel(report)
        report_id = report.id
        self.db.rollback()
        self.assertEqual(self.rows(ExcelRecord), [])
        with self.sessions() as db:
            ToolExecutionService(db, self.settings).execute("EXCEL_REPORT", report_id=report_id)
        workbook = load_workbook(self.settings.excel_path)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.max_row, 2)
        self.assertEqual(len(self.rows(ExcelRecord)), 1)

    def test_delivery_retries_then_succeeds_with_audit_per_attempt(self):
        self.settings.alert_email_delivery_mode = "smtp"
        self.settings.smtp_host = "test.invalid"
        self.settings.alert_email_from = "sender@example.test"
        self.settings.alert_email_to = "recipient@example.test"
        jobs = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)
        self.drain(jobs[:2])
        with patch.object(ToolOrchestrationService, "_send_alert_email", side_effect=[OSError("injected"), None]) as send:
            self.worker._run_job(jobs[2].id)
            self.assertEqual(self.rows(ToolJob)[2].status, "PENDING")
            self.retry_now(jobs[2].id)
            self.worker._run_job(jobs[2].id)
            self.worker._run_job(jobs[2].id)
            self.assertEqual(send.call_count, 2)
        self.assertEqual(self.rows(ToolJob)[2].status, "SUCCESS")
        self.assertEqual([row.status for row in self.rows(AlertRecord)], ["FAILED", "SUCCESS"])
        self.assertEqual([row.status for row in self.rows(ToolAuditRecord)][-2:], ["FAILED", "SUCCESS"])

    def test_dead_dependency_terminates_alert_without_sending(self):
        jobs = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)
        jobs[1].status = "DEAD"
        self.db.commit()
        self.worker._run_job(jobs[2].id)
        self.assertEqual(self.rows(ToolJob)[2].status, "DEAD")
        self.assertEqual(self.rows(AlertRecord), [])
        self.assertEqual(len(self.rows(DeadLetterRecord)), 1)

    def test_ledger_failure_does_not_block_independent_case_alert(self):
        jobs = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)
        with patch.object(ToolOrchestrationService, "write_excel", side_effect=OSError("injected disk failure")):
            self.drain(jobs)
        self.assertEqual([job.status for job in self.rows(ToolJob)], ["PENDING", "SUCCESS", "SUCCESS"])
        self.assertEqual(self.rows(ExcelRecord), [])
        self.assertEqual(self.rows(AlertRecord)[0].status, "SUCCESS")

    def test_max_attempts_create_one_dead_letter_and_do_not_rerun(self):
        jobs = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)
        jobs[1].max_attempts = 1
        self.db.commit()
        with patch.object(ToolOrchestrationService, "create_case", side_effect=RuntimeError("injected")):
            self.worker._run_job(jobs[1].id)
            self.worker._run_job(jobs[1].id)
        self.assertEqual(self.rows(ToolJob)[1].status, "DEAD")
        self.assertEqual(len(self.rows(DeadLetterRecord)), 1)
        self.assertEqual(self.rows(ToolAuditRecord)[0].status, "FAILED")

    def test_transport_failure_before_tool_entry_still_has_failed_audit(self):
        self.settings.tool_queue_enabled = False
        job = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)[0]
        with patch.object(self.worker, "_execute", side_effect=McpToolError("injected transport failure")):
            self.worker._run_job(job.id)
        audit = self.rows(ToolAuditRecord)[0]
        self.assertEqual(audit.status, "FAILED")
        self.assertEqual(json.loads(audit.payload)["source"], "mcp-dispatch")
        self.assertEqual(self.rows(ToolJob)[0].status, "PENDING")

    def test_lost_mcp_reply_after_last_attempt_does_not_mark_success_as_dead(self):
        job = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)[0]
        job.max_attempts = 1
        self.db.commit()

        def lost_reply(db, task):
            ToolExecutionService(db, self.settings).execute(
                task.kind, report_id=task.report_id, job_id=task.id, claim_token=task.claim_token,
            )
            raise McpToolError("injected lost success reply")

        with patch.object(self.worker, "_execute", lost_reply):
            self.worker._run_job(job.id)
        self.assertEqual(self.rows(ToolJob)[0].status, "SUCCESS")
        self.assertEqual(self.rows(DeadLetterRecord), [])


class GovernedExecutionTests(ToolRuntimeFixture):
    def test_illegal_low_risk_and_unknown_jobs_are_blocked_without_side_effects(self):
        report = self.report("LOW")
        for kind in ["CASE_CREATE", "ALERT_SEND", "DELETE_EVERYTHING"]:
            job = ToolJob(report_id=report.id, kind=kind, status="PENDING")
            self.db.add(job)
            self.db.commit()
            self.worker._run_job(job.id)
        self.assertEqual({row.status for row in self.rows(ToolJob)}, {"DEAD"})
        self.assertTrue(all(row.status == "BLOCKED" and not row.allowed for row in self.rows(ToolAuditRecord)))
        self.assertEqual(self.rows(RiskCase), [])
        self.assertEqual(self.rows(AlertRecord), [])

    def test_valid_jobs_commit_audits_linked_to_job_and_claim(self):
        jobs = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)
        self.drain(jobs)
        audits = self.rows(ToolAuditRecord)
        self.assertEqual(len(audits), 3)
        self.assertEqual({row.status for row in audits}, {"SUCCESS"})
        self.assertEqual({row.job_id for row in audits}, {row.id for row in jobs})
        self.assertTrue(all(json.loads(row.payload)["claimToken"] for row in audits))

    def test_stale_claim_is_blocked_at_execution_boundary(self):
        report = self.report()
        job = ToolQueueService(self.db, self.settings).enqueue_report(report.id)[0]
        with self.assertRaises(ToolAuthorizationError):
            ToolExecutionService(self.db, self.settings).execute(
                job.kind, report_id=report.id, job_id=job.id, claim_token="forged-token",
            )
        self.assertEqual(self.rows(ExcelRecord), [])
        self.assertFalse(self.rows(ToolAuditRecord)[0].allowed)

    def test_case_followup_requires_known_admin_and_is_audited(self):
        report = self.report()
        executor = ToolExecutionService(self.db, self.settings)
        executor.execute("CASE_CREATE", report_id=report.id)
        case = self.rows(RiskCase)[0]
        for actor in ["student", "invented-admin"]:
            with self.assertRaises(ToolAuthorizationError):
                executor.execute("CASE_ACK", case_id=case.id, actor=actor)
        self.assertEqual(self.rows(CaseNote), [])
        executor.execute("CASE_ACK", case_id=case.id, actor="admin")
        executor.execute("CASE_NOTE_ADD", case_id=case.id, actor="admin", note="synthetic follow-up")
        self.assertEqual(self.rows(RiskCase)[0].status, "ACKNOWLEDGED")
        self.assertEqual(len(self.rows(CaseNote)), 2)
        self.assertEqual([row.status for row in self.rows(ToolAuditRecord)], ["SUCCESS", "BLOCKED", "BLOCKED", "SUCCESS", "SUCCESS"])

    def test_client_rejects_business_failure_even_without_mcp_error_flag(self):
        async def failed_call(*args, **kwargs):
            return SimpleNamespace(isError=False, content=[SimpleNamespace(text="FAILED: no SMTP")])
        client = MindBridgeMcpToolClient(self.settings)
        with self.assertRaises(McpToolError):
            asyncio.run(client._call_tool(SimpleNamespace(call_tool=failed_call), "mindbridge_alert_send", {}))


class LegacySchemaTests(unittest.TestCase):
    def test_additive_migration_preserves_duplicate_history_and_is_repeatable(self):
        engine = create_engine("sqlite://")
        self.addCleanup(engine.dispose)
        with engine.begin() as connection:
            connection.exec_driver_sql("CREATE TABLE tool_jobs (id INTEGER PRIMARY KEY, report_id INTEGER, kind VARCHAR(32), status VARCHAR(32))")
            connection.exec_driver_sql("INSERT INTO tool_jobs VALUES (1, 7, 'EXCEL_REPORT', 'SUCCESS'), (2, 7, 'EXCEL_REPORT', 'DEAD')")
        upgrade_tool_jobs(engine)
        upgrade_tool_jobs(engine)
        self.assertTrue({"idempotency_key", "claim_token", "lease_expires_at", "correlation_json"}.issubset(
            {column["name"] for column in inspect(engine).get_columns("tool_jobs")}
        ))
        with engine.connect() as connection:
            rows = connection.execute(text("SELECT id, idempotency_key FROM tool_jobs ORDER BY id")).all()
        self.assertEqual(rows, [(1, "7:EXCEL_REPORT"), (2, None)])
        with self.assertRaises(IntegrityError), engine.begin() as connection:
            connection.exec_driver_sql("INSERT INTO tool_jobs (id,report_id,kind,status,idempotency_key) VALUES (3,7,'EXCEL_REPORT','PENDING','7:EXCEL_REPORT')")


class McpTransportTests(ToolRuntimeFixture):
    def test_real_stdio_denies_low_risk_and_missing_report(self):
        report_id = self.report("LOW").id

        async def invoke():
            async with MindBridgeMcpToolClient(self.settings)._session() as session:
                low = await session.call_tool("mindbridge_case_create", {"report_id": report_id})
                missing = await session.call_tool("mindbridge_excel_report", {"report_id": 99999})
                self.assertTrue(low.isError)
                self.assertTrue(missing.isError)

        asyncio.run(invoke())
        self.assertEqual(self.rows(RiskCase), [])
        audits = self.rows(ToolAuditRecord)
        self.assertEqual(len(audits), 2)
        self.assertTrue(all(not row.allowed and row.status == "BLOCKED" for row in audits))
        self.assertTrue(all(json.loads(row.payload)["source"] == "mcp-stdio" for row in audits))

    def test_durable_jobs_execute_over_real_stdio_and_remain_idempotent(self):
        self.settings.tool_queue_enabled = False
        jobs = ToolQueueService(self.db, self.settings).enqueue_report(self.report().id)
        self.drain(jobs)
        self.drain(jobs)
        self.assertEqual({job.status for job in self.rows(ToolJob)}, {"SUCCESS"})
        self.assertEqual(len(self.rows(ExcelRecord)), 1)
        self.assertEqual(len(self.rows(RiskCase)), 1)
        self.assertEqual(len(self.rows(AlertRecord)), 1)
        audits = self.rows(ToolAuditRecord)
        self.assertEqual(len(audits), 3)
        self.assertTrue(all(row.status == "SUCCESS" and row.job_id is not None for row in audits))
        self.assertTrue(all(json.loads(row.payload)["source"] == "mcp-stdio" for row in audits))

    def test_real_stdio_delivery_failure_is_error_and_keeps_failed_audit(self):
        report = self.report()
        ToolExecutionService(self.db, self.settings).execute("CASE_CREATE", report_id=report.id)
        case_id = self.rows(RiskCase)[0].id
        self.settings.alert_email_delivery_mode = "smtp"
        self.settings.smtp_host = ""

        async def invoke():
            async with MindBridgeMcpToolClient(self.settings)._session() as session:
                result = await session.call_tool("mindbridge_alert_send", {"case_id": case_id})
                self.assertTrue(result.isError)

        asyncio.run(invoke())
        self.assertEqual(self.rows(AlertRecord)[0].status, "FAILED")
        self.assertEqual(self.rows(ToolAuditRecord)[-1].status, "FAILED")


if __name__ == "__main__":
    unittest.main()
