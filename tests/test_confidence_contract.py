import json
import unittest
from dataclasses import replace
from types import SimpleNamespace
from unittest.mock import Mock, patch

from sqlalchemy import create_engine, inspect
from sqlalchemy.exc import IntegrityError
from openpyxl import load_workbook

from app.agents.autonomous import UnderstandingAgent, SafetyAgent, ContextAgent, ResponseAgent
from app.agents.coordinator import EventDrivenCoordinator
from app.agents.events import AgentArtifact, AgentEventType, AgentTask, AgentTurnResult
from app.agents.event_driven_runtime import EventDrivenAgentRuntimeService
from app.agents.harness import MindBridgeAgentHarness
from app.agents.plan_validation import validate_plan
from app.agents.registry import AgentDecision, AgentRegistry
from app.core.confidence_migration import upgrade_confidence_storage
from app.core.config import Settings
from app.core.enums import RiskLevel
from app.models.entities import PsychologicalReport, ToolJob
from app.schemas.dtos import AiMessage, ChatRequest
from app.services.ai import AiClient
from app.services.assessment import PsychologicalAssessmentService
from app.services.confidence import confidence_info, confidence_summary, report_confidence_info
from test_tool_reliability import ToolRuntimeFixture


class ConfidenceMeaningTests(unittest.TestCase):
    def assess(self, confidence, include=True):
        data = dict(emotion='HIGH_RISK', emotionScore=4, risk='HIGH', summary='synthetic model assessment')
        if include: data['confidence'] = confidence
        client = Mock(); client.complete.return_value = json.dumps(data)
        return PsychologicalAssessmentService(client).assess('我需要认真聊聊最近的情况。')

    def test_rule_risk_has_no_fabricated_probability(self):
        client = Mock(); client.complete.side_effect = AssertionError('rules should decide')
        result = PsychologicalAssessmentService(client).assess('我今晚想伤害自己。')
        self.assertEqual(result.risk, RiskLevel.HIGH)
        self.assertIsNone(result.confidence)
        self.assertEqual(result.confidence_info['kind'], 'not_applicable')
        self.assertFalse(result.confidence_info['usedForRouting'])
        client.complete.assert_not_called()

    def test_low_or_high_self_report_does_not_change_danger_routing(self):
        for value in [0., .01, .99, 1.]:
            with self.subTest(value=value):
                result = self.assess(value)
                self.assertEqual(result.risk, RiskLevel.HIGH)
                self.assertEqual(result.confidence, value)
                self.assertEqual(result.confidence_info['calibrationStatus'], 'not_calibrated')

    def test_invalid_optional_confidence_cannot_erase_valid_high_assessment(self):
        for value in [None, -1, 2, float('nan'), float('inf'), '0.9', True, {}, []]:
            with self.subTest(value=value):
                result = self.assess(value)
                self.assertEqual(result.risk, RiskLevel.HIGH)
                self.assertEqual(result.source, 'model')
                self.assertFalse(result.degraded)
                self.assertIsNone(result.confidence)
                self.assertIn(result.confidence_info['validation'], {'missing', 'invalid'})
        self.assertEqual(self.assess(None, include=False).confidence_info['validation'], 'missing')

    def test_legacy_value_is_preserved_without_inventing_provenance(self):
        report = SimpleNamespace(confidence=.95, confidence_metadata_json=None)
        info = report_confidence_info(report)
        self.assertEqual(info['kind'], 'legacy_unknown')
        self.assertIsNone(info['value']); self.assertEqual(info['legacyValue'], .95)
        self.assertIn('来源未标注', confidence_summary(report))

    def test_display_reconstructs_canonical_uncalibrated_metadata(self):
        metadata = confidence_info(.8, 'model', 'valid'); metadata['calibrationStatus'] = 'calibrated'
        report = SimpleNamespace(confidence=.8, confidence_metadata_json=json.dumps(metadata))
        self.assertEqual(report_confidence_info(report)['calibrationStatus'], 'not_calibrated')
        self.assertIn('未校准', confidence_summary(report))
        report.confidence_metadata_json = '{"version":"confidence-semantics-v1"}'
        self.assertEqual(report_confidence_info(report)['kind'], 'legacy_unknown')

    def test_scheduler_priority_is_bounded_preference_and_artifacts_have_no_score(self):
        self.assertEqual(AgentDecision(True, .8).claim_priority, .8)
        for value in [True, -1, 2, float('nan'), float('inf')]:
            with self.subTest(value=value), self.assertRaises(ValueError): AgentDecision(True, value)
        artifact = AgentArtifact(id='x', owner='ResponseAgent', kind='response_proposal', payload={})
        self.assertFalse(hasattr(artifact, 'confidence'))


class PlanContractTests(ToolRuntimeFixture):
    def board(self, text='最近有点压力，想聊聊。', empty_retrieval=False):
        runtime = EventDrivenAgentRuntimeService(self.db, self.settings)
        services, board = runtime.prepare_turn(self.user, self.session, text, text)
        if empty_retrieval: services.knowledge.retrieve = lambda *args, **kwargs: []
        for kind, agent_type in [('understanding', UnderstandingAgent), ('risk', SafetyAgent),
                                 ('context', ContextAgent), ('response', ResponseAgent), ('plan_review', SafetyAgent)]:
            task = AgentTask(id=kind, title=kind, metadata={'kind': kind})
            agent = agent_type(services)
            board = board.add_task(task).apply_turn_result(task, agent.name, agent.act(task, board))
        return board

    def alter(self, board, kind, *, payload=None, **changes):
        original = board.latest_artifact(kind)
        return board.add_artifact(replace(original, payload=payload if payload is not None else original.payload, **changes))

    def test_real_agent_preparation_passes_chat_consult_and_risk(self):
        for text in ['解释Python字典。', '最近有点压力，想聊聊。', '我今晚想伤害自己。']:
            with self.subTest(text=text):
                decision = validate_plan(self.board(text))
                self.assertTrue(decision.approved, decision)
                self.assertFalse(decision.confidenceUsed)
                self.assertEqual(decision.scope, 'prompt_plan_structure_only')

    def test_no_retrieved_evidence_allows_declared_generic_crisis_support(self):
        decision = validate_plan(self.board('我今晚想伤害自己。', empty_retrieval=True))
        self.assertTrue(decision.approved, decision)
        self.assertEqual(decision.evidenceMode, 'generic_support')

    def test_each_preparation_requirement_rejects_corrupt_candidate(self):
        mutations = [
            ('currentArtifactBindings', lambda p: p.update(inputArtifactIds={})),
            ('routeMatchesAssessment', lambda p: p.update(risk='HIGH')),
            ('routeMatchesAssessment', lambda p: p.update(mode='normal_chat')),
            ('messagesAndPolicyPrepared', lambda p: p.update(messages=[])),
            ('messagesAndPolicyPrepared', lambda p: p.update(messages=[AiMessage(role='user', content='当前安全 可信任的人 紧急')])),
            ('messagesAndPolicyPrepared', lambda p: p.update(messages=[*p['messages'][:-1], AiMessage(role='user', content='stale input')])),
            ('messagesAndPolicyPrepared', lambda p: p.update(safetyPolicy={})),
            ('evidenceModeDeclared', lambda p: p.update(evidenceMode='general_chat')),
        ]
        for failure, mutate in mutations:
            with self.subTest(failure=failure):
                board = self.board(); payload = dict(board.latest_artifact('response_proposal').payload); mutate(payload)
                decision = validate_plan(self.alter(board, 'response_proposal', payload=payload))
                self.assertFalse(decision.approved); self.assertIn(failure, decision.violations)

    def test_review_requires_independent_owner_binding_scope_and_exact_boolean(self):
        for change in ['owner', 'binding', 'scope', 'boolean', 'checks']:
            with self.subTest(change=change):
                board = self.board(); review = board.latest_artifact('plan_review'); payload = dict(review.payload)
                if change == 'owner': board = self.alter(board, 'plan_review', owner='ResponseAgent')
                else:
                    if change == 'binding': payload['responseArtifactId'] = 'old-response'
                    if change == 'scope': payload['scope'] = 'final_reply'
                    if change == 'boolean': payload['approved'] = 'true'
                    if change == 'checks': payload['checks'] = {}
                    board = self.alter(board, 'plan_review', payload=payload)
                self.assertFalse(validate_plan(board).approved)

    def test_uncertainty_requires_server_clarification_message(self):
        board = self.board('我目前不敢保证自己是安全的。')
        self.assertTrue(validate_plan(board).approved)
        response = board.latest_artifact('response_proposal')
        messages = [m for m in response.payload['messages'] if '本轮存在未能确认的安全信号' not in m.content]
        decision = validate_plan(self.alter(board, 'response_proposal', payload={**response.payload, 'messages': messages}))
        self.assertFalse(decision.approved); self.assertIn('clarificationPrepared', decision.violations)

    def test_retired_threshold_does_not_block_valid_plan_and_rejection_is_audited(self):
        board = self.board(); settings = self.settings.model_copy(update={'agent_final_acceptance_min_confidence': 1.})
        coordinator = EventDrivenCoordinator(AgentRegistry([]), SimpleNamespace(name='CoordinatorAgent', remember_acceptance=lambda *a: None), settings)
        accepted = coordinator._try_accept_final(board)
        self.assertTrue(accepted.final_artifact_id)
        event = accepted.events[-1]
        self.assertEqual(event.type, AgentEventType.PLAN_ACCEPTED)
        self.assertFalse(event.metadata['confidenceUsed'])
        response = board.latest_artifact('response_proposal')
        broken = self.alter(board, 'response_proposal', payload={**response.payload, 'safetyPolicy': {}})
        rejected = coordinator._try_accept_final(broken)
        rejected = coordinator._try_accept_final(rejected)
        self.assertFalse(rejected.final_artifact_id)
        self.assertEqual(sum(e.type == AgentEventType.PLAN_REJECTED for e in rejected.events), 1)

    def test_invalid_plan_keeps_high_report_and_outbox_but_never_approves_plan(self):
        original = ResponseAgent.act
        def invalid(agent, task, board):
            result = original(agent, task, board)
            artifacts = tuple(replace(a, payload={**a.payload, 'inputArtifactIds': {}}) for a in result.artifacts)
            return replace(result, artifacts=artifacts)
        with patch.object(ResponseAgent, 'act', invalid):
            outcome = MindBridgeAgentHarness(self.db, self.settings).run(self.user, ChatRequest(message='我今晚想伤害自己。'))
        self.assertFalse(outcome.plan_accepted)
        self.assertEqual(outcome.risk_level, 'HIGH')
        self.assertEqual(self.db.query(ToolJob).filter_by(report_id=outcome.report_id).count(), 3)

    def test_low_self_reported_confidence_still_persists_high_risk_actions(self):
        original = AiClient.complete
        def complete(ai, messages, **kwargs):
            if messages[0].content.startswith('你负责分析校园心理健康消息'):
                return '{"emotion":"HIGH_RISK","emotionScore":4,"risk":"HIGH","confidence":0.01,"summary":"fixture"}'
            return original(ai, messages, **kwargs)
        with patch.object(AiClient, 'complete', complete):
            outcome = MindBridgeAgentHarness(self.db, self.settings).run(self.user, ChatRequest(message='我需要认真聊聊最近的情况。'))
        self.assertTrue(outcome.plan_accepted); self.assertEqual(outcome.risk_level, 'HIGH')
        report = self.db.get(PsychologicalReport, outcome.report_id)
        self.assertEqual(report.confidence, .01)
        self.assertEqual(self.db.query(ToolJob).filter_by(report_id=report.id).count(), 3)

    def test_rule_report_and_excel_do_not_invent_probability(self):
        outcome = MindBridgeAgentHarness(self.db, self.settings).run(self.user, ChatRequest(message='我今晚想伤害自己。'))
        jobs = self.db.query(ToolJob).filter_by(report_id=outcome.report_id).all(); self.drain(jobs)
        report = self.db.get(PsychologicalReport, outcome.report_id)
        self.assertIsNone(report.confidence)
        self.assertEqual(report_confidence_info(report)['kind'], 'not_applicable')
        workbook = load_workbook(self.settings.excel_path); self.addCleanup(workbook.close)
        sheet = workbook.active; headers = [c.value for c in sheet[1]]
        row = next(r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] == report.id)
        self.assertIsNone(row[headers.index('confidence')])
        self.assertEqual(row[headers.index('confidenceKind')], 'not_applicable')


class ConfidenceMigrationTests(unittest.TestCase):
    def test_sqlite_rebuild_preserves_rows_foreign_keys_indexes_triggers_and_sequence(self):
        engine = create_engine('sqlite://'); self.addCleanup(engine.dispose)
        with engine.connect() as c:
            c.exec_driver_sql('PRAGMA foreign_keys=ON'); c.commit()
            c.exec_driver_sql('CREATE TABLE psychological_reports(id INTEGER PRIMARY KEY AUTOINCREMENT, confidence FLOAT NOT NULL, content TEXT UNIQUE)')
            c.exec_driver_sql('CREATE INDEX confidence_index ON psychological_reports(confidence)')
            c.exec_driver_sql('CREATE TABLE children(report_id INTEGER REFERENCES psychological_reports(id))')
            c.exec_driver_sql('CREATE TABLE writes(id INTEGER)')
            c.exec_driver_sql('CREATE TRIGGER report_insert AFTER INSERT ON psychological_reports BEGIN INSERT INTO writes VALUES(new.id); END')
            c.exec_driver_sql("INSERT INTO psychological_reports VALUES(1,.95,'old'),(2,.8,'deleted')")
            c.exec_driver_sql('INSERT INTO children VALUES(1)'); c.exec_driver_sql('DELETE FROM psychological_reports WHERE id=2'); c.commit()
        upgrade_confidence_storage(engine); upgrade_confidence_storage(engine)
        self.assertTrue(next(c for c in inspect(engine).get_columns('psychological_reports') if c['name']=='confidence')['nullable'])
        with engine.connect() as c:
            self.assertEqual(c.exec_driver_sql('SELECT confidence,content,confidence_metadata_json FROM psychological_reports').all(),[(.95,'old',None)])
            self.assertEqual(c.exec_driver_sql('PRAGMA foreign_keys').scalar(), 1)
            self.assertEqual(c.exec_driver_sql('PRAGMA foreign_key_check').all(), [])
            c.exec_driver_sql("INSERT INTO psychological_reports(confidence,content) VALUES(NULL,'new')"); c.commit()
            self.assertEqual(c.exec_driver_sql("SELECT id FROM psychological_reports WHERE content='new'").scalar(), 3)
            self.assertEqual(c.exec_driver_sql('SELECT count(*) FROM writes').scalar(), 3)
            self.assertIn('confidence_index', {i['name'] for i in inspect(c).get_indexes('psychological_reports')})

    def test_nullable_legacy_schema_only_adds_missing_metadata(self):
        engine = create_engine('sqlite://'); self.addCleanup(engine.dispose)
        with engine.begin() as c:
            c.exec_driver_sql('CREATE TABLE psychological_reports(id INTEGER PRIMARY KEY, confidence FLOAT)')
            c.exec_driver_sql('INSERT INTO psychological_reports VALUES(1,NULL)')
        upgrade_confidence_storage(engine); upgrade_confidence_storage(engine)
        with engine.connect() as c:
            self.assertEqual(c.exec_driver_sql('SELECT * FROM psychological_reports').all(), [(1,None,None)])
