"""Read physical evaluation evidence, never infer it from tool success strings."""
from __future__ import annotations

from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import threading
import time


class EvidenceUnavailable(RuntimeError):
    pass


def read_ledger(path):
    from openpyxl import load_workbook
    path=Path(path)
    if not path.exists(): return {'counts':{},'sha256':None,'rows':[]}
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        rows=list(book.active.iter_rows(values_only=True))
        if not rows or 'reportId' not in rows[0]: raise EvidenceUnavailable('Ledger header is not verifiable')
        index=rows[0].index('reportId');ids=[str(row[index]) for row in rows[1:] if row and row[index] is not None]
        return {'counts':dict(Counter(ids)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'rows':ids}
    finally:book.close()


def read_mailpit(base_url, recipient, subject_prefix):
    import httpx
    if base_url!='http://127.0.0.1:8025': raise ValueError('Only the local Mailpit sink is permitted')
    records=[];start=0
    with httpx.Client(trust_env=False,timeout=5) as client:
        while True:
            r=client.get(base_url+'/api/v1/search',params={'query':'to:'+recipient,'start':start,'limit':100})
            r.raise_for_status();page=r.json()
            if not isinstance(page.get('messages'),list): raise EvidenceUnavailable('Mailpit response missing messages')
            for row in page['messages']:
                addresses=[x.get('Address','').lower() for x in row.get('To',[])]
                if recipient.lower() not in addresses or not row.get('Subject','').startswith(subject_prefix):
                    raise EvidenceUnavailable('Mailpit namespace mismatch')
                m=re.search(r'\breportId=(\d+)\b',row['Subject'])
                if not m: raise EvidenceUnavailable('Mailpit message cannot be bound to report')
                records.append({'id':row['ID'],'reportId':m[1],'subject':row['Subject']})
            start+=len(page['messages'])
            # Mailpit total is the whole mailbox, not this search result.
            matched=page.get('messages_count')
            if not isinstance(matched,int):raise EvidenceUnavailable('Mailpit matched-count unavailable')
            if start>=matched:break
            if not page['messages']:raise EvidenceUnavailable('Mailpit pagination incomplete')
    if len({r['id'] for r in records})!=len(records):raise EvidenceUnavailable('Mailpit pagination repeated a message')
    return {'counts':dict(Counter(r['reportId'] for r in records)),'messages':records}


def identify_reply(text,outcome,*,fixed_template,disconnected=False,access_denied=False,has_error=False):
    sha=hashlib.sha256(text.encode()).hexdigest() if text else None
    if disconnected and not text:return 'none',sha
    if outcome and outcome.get('type')=='FINAL_ACCEPTED' and outcome.get('finalSha256')==sha:return 'reviewed_reply',sha
    if outcome and outcome.get('type')=='SAFE_FALLBACK' and outcome.get('finalSha256')==sha:return 'safe_fallback',sha
    # HTTP deadline fallback may have no persisted body-review outcome.
    # Compare the actual received bytes to the server's fixed template.
    if text and text==fixed_template:return 'safe_fallback',sha
    if access_denied and has_error and not text:return 'access_error',sha
    return ('error' if has_error else 'none'),sha


class EventRecorder:
    def __init__(self,task):
        self.task=task;self.lock=threading.RLock();self.seq=0;self.active=None
        self.turns={};self.calls=[];self.raw=[];self.thread=threading.local();self.embedding_calls=0
        self.faults={f['id']:{'id':f['id'],'stage':f['stage'],'trigger':f['trigger'],'triggered':False}
                     for f in task['environment']['faults']}

    def begin(self,turn):
        self.active=turn['id'];self.turns[self.active]={'events':[],'runId':None,'route':{},'context':{},
            'planChecks':{},'traceIds':[],'sse':[],'started':time.perf_counter(),'firstTextMs':None,'bodyCaptured':False}

    def tick(self):
        with self.lock:self.seq+=1;return self.seq

    def event(self,kind,**fields):
        with self.lock:
            if self.active is None:return
            turn=self.turns[self.active]
            fields.setdefault('runId',turn['runId'] or 'unavailable')
            row={'seq':self.tick(),'kind':kind,**fields};turn['events'].append(row)
            self.raw.append({'event':kind,'seq':row['seq'],'turnId':self.active,'time':time.perf_counter()})
            return row

    def trigger(self,kind,stage=None):
        for f in self.task['environment']['faults']:
            if f['kind']==kind and (stage is None or f['stage']==stage):
                self.faults[f['id']]['triggered']=True
                self.raw.append({'event':'FAULT_TRIGGERED','id':f['id'],'stage':f['stage'],'trigger':f['trigger'],'seq':self.tick()})
                return f
        return None

    def has(self,kind,stage=None):
        return next((f for f in self.task['environment']['faults'] if f['kind']==kind and (stage is None or f['stage']==stage)),None)

    def summary(self):
        return {'events':self.raw,'invocations':self.calls,'faults':list(self.faults.values()),
                'embeddingRequests':self.embedding_calls,'source':'observed_test_process_not_model_claims'}


def controlled_embedding(text):
    """Declared lexical hash embedding for plumbing only; never called Qwen."""
    import math
    values=[0.0]*1024
    for token in re.findall(r'[a-zA-Z0-9_]+|[\u4e00-\u9fff]',text.lower()):
        bucket=int(hashlib.sha256(token.encode()).hexdigest()[:8],16)%len(values);values[bucket]+=1
    norm=math.sqrt(sum(x*x for x in values)) or 1
    return [x/norm for x in values]
