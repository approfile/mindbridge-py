"""Verify a running local stack using HTTP and physical side-effect evidence.

Run inside the app container so database/Redis settings match its environment.
Only synthetic conversations are created; existing rows are never deleted.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import sys
import tempfile
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import httpx
from openpyxl import load_workbook
import redis
from sqlalchemy import text

from app.core.config import get_settings
from app.core.database import SessionLocal, engine
from app.models.entities import AgentRunTrace, ChatMessage, ChatSession, KnowledgeChunk, PsychologicalReport, ToolJob
from app.services.knowledge import KnowledgeService
from app.services.memory import RedisShortTermMemoryStore
from app.services.privacy_views import RESTRICTED


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base-url', default='http://127.0.0.1:8080')
    parser.add_argument('--mailpit-url', default='http://mailpit:8025')
    parser.add_argument('--expect-provider', choices=['mock', 'openai', 'ollama'], default='mock')
    args = parser.parse_args()
    settings = get_settings()
    stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    output = ROOT / 'target' / 'local-verification' / stamp
    output.mkdir(parents=True)
    report = {'createdAt': stamp, 'provider': args.expect_provider, 'checks': [], 'passed': False}

    def check(name, evidence):
        report['checks'].append({'name': name, 'passed': True, 'evidence': evidence})
        print(f'PASS {name}', flush=True)

    try:
        with httpx.Client(base_url=args.base_url, timeout=90, trust_env=False) as client:
            student = ('student', 'student123')
            admin = ('admin', 'admin123')
            def get(path, auth=admin):
                response = client.get(path, auth=auth)
                response.raise_for_status()
                return response.json()

            require(get('/actuator/health', None)['status'] == 'UP', 'health')
            for page in ['/', '/student.html', '/admin.html', '/student.js', '/admin.js', '/styles.css']:
                require(client.get(page).status_code == 200, f'frontend {page}')
            require(client.get('/api/profile').status_code == 401, 'unauthenticated access')
            require(get('/api/profile', student)['username'] == 'student', 'student auth')
            require(get('/api/profile')['username'] == 'admin', 'admin auth')
            require(client.get('/api/admin/reports', auth=student).status_code == 403, 'student admin denial')
            status = get('/api/agent/status', student)
            require(status['provider'] == args.expect_provider, 'unexpected model provider')
            require(len(status['agents']) == 5 and len(status['skills']) == 7, 'agent/skill registration')
            check('HTTP, frontend, roles, agents and skills', {'agents': 5, 'skills': 7})

            cache = redis.Redis.from_url(settings.redis_url, decode_responses=True, socket_timeout=3)
            require(cache.ping(), 'Redis unavailable')
            with engine.connect() as conn:
                require(conn.execute(text('SELECT 1')).scalar() == 1, 'SQL unavailable')
            check('real SQL and Redis connections', {'database': engine.dialect.name, 'redis': 'PONG'})

            # This checks the installed native vector backend only. Synthetic
            # vectors are never presented as a real embedding/model result.
            import chromadb
            from chromadb.config import Settings as ChromaSettings
            with tempfile.TemporaryDirectory(prefix='mindbridge-chroma-probe-') as folder:
                vector_client = chromadb.PersistentClient(path=folder, settings=ChromaSettings(anonymized_telemetry=False))
                collection = vector_client.create_collection('local-stack-probe', embedding_function=None)
                collection.upsert(ids=['probe'], embeddings=[[1.0, 0.0, 0.0]], documents=['synthetic backend probe'])
                result = collection.query(query_embeddings=[[1.0, 0.0, 0.0]], n_results=1)
                require(result['ids'] == [['probe']], 'Chroma native backend failed')
                require(Path(folder, 'chroma.sqlite3').exists(), 'Chroma persistence missing')
                vector_client.delete_collection('local-stack-probe')
            check('Chroma native storage and query', {'embedding': 'synthetic storage probe; no semantic embedding claim'})

            sessions = {}
            expected_reports = {}
            reply_outcomes = []
            cases = [
                ('chat', '你好，今天校园里的天气不错。', None),
                ('consult', '我最近考试压力很大，连续几天失眠，想找人聊聊。', 'LOW'),
                ('risk', '我不想活了，想结束生命。', 'HIGH'),
            ]
            for name, message, risk in cases:
                response = client.post('/api/chat/stream', auth=student, json={'message': message})
                response.raise_for_status()
                events = []
                for block in response.text.strip().split('\n\n'):
                    values = dict(line.split(': ', 1) for line in block.splitlines() if ': ' in line)
                    if 'event' in values and 'data' in values:
                        events.append((values['event'], json.loads(values['data'])))
                require(any(name == 'done' for name, _ in events), f'{name} SSE missing done')
                require(any(name == 'token' and data.get('content') for name, data in events), f'{name} empty reply')
                require(not any(name == 'error' for name, _ in events), f'{name} SSE error')
                session_id = next(data['sessionId'] for name, data in events if name == 'meta')
                sessions[name] = session_id
                with SessionLocal() as db:
                    session = db.query(ChatSession).filter_by(public_id=session_id).one()
                    rows = db.query(PsychologicalReport).filter_by(session_id=session.id).all()
                    if risk:
                        require(len(rows) == 1, f'{name} missing report')
                        if name == 'risk':
                            require(rows[0].risk_level == 'HIGH', 'risk override missing')
                        expected_reports[name] = rows[0].id
                    else:
                        require(not rows, 'CHAT unexpectedly generated a report')
                    trace = db.query(AgentRunTrace).filter_by(session_id=session.id).one()
                    steps = json.loads(trace.agent_steps_json)
                    require(bool(steps), 'missing runtime trace')
                    outcome = next((step for step in reversed(steps) if step.get('kind') == 'response_outcome'), None)
                    require(outcome is not None, 'missing reviewed reply outcome')
                    reply_outcomes.append({'case': name, 'type': outcome['type'], 'status': outcome['status']})
                memory = RedisShortTermMemoryStore(settings)
                require(len(memory.load_recent(session_id)) >= 2, 'real Redis memory missing')
                raw = cache.lrange(memory._key(session_id), 0, -1)
                require(raw and all(json.loads(value)['content'].startswith('mbenc:v1:') for value in raw), 'Redis plaintext')
                check(f'{name} SSE, trace, report policy and encrypted Redis',
                      {'sessionId': session_id, 'replyOutcome': reply_outcomes[-1]})

            report['replyOutcomes'] = reply_outcomes
            if args.expect_provider != 'mock':
                require(any(item['type'] == 'FINAL_ACCEPTED' for item in reply_outcomes),
                        'all real-model replies used fallback; generated reply path not verified')

            response = client.post('/api/chat/stream', auth=student,
                                   json={'message': '请接着聊聊校园生活。', 'sessionId': sessions['chat']})
            response.raise_for_status()
            require('event: done' in response.text and 'event: error' not in response.text, 'follow-up SSE failed')
            with SessionLocal() as db:
                session = db.query(ChatSession).filter_by(public_id=sessions['chat']).one()
                require(db.query(ChatMessage).filter_by(session_id=session.id).count() == 4, 'follow-up session history missing')
            require(len(RedisShortTermMemoryStore(settings).load_recent(sessions['chat'])) == 4, 'Redis follow-up history missing')
            check('same-session follow-up persistence and Redis history', {'messages': 4})

            ids = list(expected_reports.values())
            deadline = time.monotonic() + 90
            while True:
                with SessionLocal() as db:
                    jobs = db.query(ToolJob).filter(ToolJob.report_id.in_(ids)).all()
                    states = [(job.kind, job.status) for job in jobs]
                if jobs and all(state == 'SUCCESS' for _, state in states):
                    break
                require(time.monotonic() < deadline, f'outbox not completed: {states}')
                time.sleep(1)
            require({'EXCEL_REPORT', 'CASE_CREATE', 'ALERT_SEND'} <= {kind for kind, _ in states}, 'missing outbox kinds')
            workbook = load_workbook(settings.excel_path, read_only=True)
            try:
                written = [row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)]
                require(all(written.count(report_id) == 1 for report_id in ids), 'Excel missing/duplicate report')
            finally:
                workbook.close()
            check('outbox and physical Excel', {'jobs': states, 'reportIds': ids})

            alerts = get('/api/admin/alerts')
            require(any(row['reportId'] == expected_reports['risk'] and row['status'] == 'SUCCESS' for row in alerts), 'SMTP alert not successful')
            mail = httpx.get(args.mailpit_url + '/api/v1/messages', timeout=10, trust_env=False)
            mail.raise_for_status()
            messages = mail.json().get('messages', [])
            marker = f"reportId={expected_reports['risk']}"
            require(any(marker in message['Subject'] for message in messages), 'Mailpit has not received this report')
            check('SMTP physically received in Mailpit', {'subjectMarker': marker})

            hidden = get('/api/admin/reports')
            require(all(row['content'] == RESTRICTED for row in hidden), 'admin exposed original input')
            response = client.get(f"/api/admin/reports/{expected_reports['risk']}?includeSensitive=true", auth=admin,
                                  headers={'X-Access-Purpose': 'case_followup'})
            require(response.status_code == 403, 'default admin sensitive read must be denied')
            for path in ['cases', 'excel-records', 'tool-jobs', 'tool-audits', 'dead-letters', 'agent-traces']:
                get('/api/admin/' + path)
            with engine.connect() as conn:
                encrypted = conn.execute(text('SELECT content FROM psychological_reports WHERE id=:id'), {'id': ids[-1]}).scalar_one()
                require(encrypted.startswith('mbenc:v1:'), 'SQL plaintext report')
            check('admin projections, sensitive permission, SQL encryption', {})

            with SessionLocal() as db:
                knowledge = KnowledgeService(db, settings)
                matches = knowledge.retrieve('考试压力很大导致失眠怎么办')
                require(bool(matches), 'knowledge retrieval empty')
                info = knowledge.status()
                require(info['databaseChunks'] >= 34, 'knowledge not seeded')
                if settings.knowledge_vector_required:
                    require(info['vectorAvailable'] and info['vectorChunks'] == info['databaseChunks'], 'vector index incomplete')
                    chunks = db.query(KnowledgeChunk).all()
                    dimensions = {len(json.loads(chunk.embedding_json)) for chunk in chunks if chunk.embedding_json}
                    require(all(chunk.embedding_json for chunk in chunks) and len(dimensions) == 1,
                            'SQL embedding persistence missing or inconsistent')
                    info['persistedEmbeddingDimensions'] = next(iter(dimensions))
                check('knowledge retrieval', info)
            report['passed'] = True
    except Exception as exc:
        # Avoid persisting provider error responses, URLs, headers or credentials.
        report['failureType'] = type(exc).__name__
        if isinstance(exc, AssertionError):
            report['failure'] = str(exc)
        print(f'FAIL {type(exc).__name__}', flush=True)
        if isinstance(exc, AssertionError):
            print(str(exc), flush=True)
    finally:
        destination = output / 'report.json'
        destination.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
        print(f'Report: {destination}', flush=True)
    return 0 if report['passed'] else 1


if __name__ == '__main__':
    raise SystemExit(main())
